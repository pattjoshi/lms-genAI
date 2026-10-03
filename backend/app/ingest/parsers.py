"""Step 1 of ingestion: turn a file into clean text sections.

A Section is a piece of text plus WHERE it came from (page number for PDFs, heading for
DOCX / HTML / TXT). Keeping the location from the very first step is what makes
citations like "ML201_M1.pdf, page 3" possible at the end of the pipeline.
"""

import re
from collections import Counter
from dataclasses import dataclass
from pathlib import Path

SUPPORTED_TYPES = {".pdf": "pdf", ".docx": "docx", ".html": "html", ".htm": "html", ".txt": "txt"}


class ParseError(Exception):
    """The file can't be read. The message is shown to the teacher."""


@dataclass
class Section:
    text: str
    page: int | None = None  # 1-based, PDFs only
    heading: str | None = None


def file_type_for(file_name: str) -> str:
    suffix = Path(file_name).suffix.lower()
    if suffix == ".doc":
        raise ParseError("Old Word .doc files aren't supported. Open the file in Word and save it as .docx.")
    if suffix not in SUPPORTED_TYPES:
        raise ParseError(f"Unsupported file type '{suffix}'. Upload PDF, DOCX, HTML or TXT.")
    return SUPPORTED_TYPES[suffix]


def parse_file(path: Path, file_type: str) -> list[Section]:
    parser = {"pdf": parse_pdf, "docx": parse_docx, "html": parse_html, "txt": parse_txt}[file_type]
    try:
        sections = parser(path)
    except ParseError:
        raise
    except Exception as exc:  # noqa: BLE001 - library errors become a readable message
        raise ParseError(f"Could not read this {file_type.upper()} file ({type(exc).__name__}: {exc}).") from exc
    sections = [s for s in sections if s.text.strip()]
    if not sections:
        hint = " It may be a scanned PDF (images only), which needs OCR." if file_type == "pdf" else ""
        raise ParseError(f"No text found in the file.{hint}")
    return sections


def _clean(text: str) -> str:
    text = text.replace(" ", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"-\n(?=[a-z])", "", text)  # re-join words hyphenated across lines
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


# ------------------------------------------------------------------------- PDF


def parse_pdf(path: Path) -> list[Section]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    if reader.is_encrypted:
        raise ParseError("This PDF is password-protected. Upload an unlocked copy.")
    pages = [page.extract_text() or "" for page in reader.pages]
    pages = strip_repeated_lines(pages)
    return [Section(text=_clean(text), page=number) for number, text in enumerate(pages, start=1)]


def strip_repeated_lines(pages: list[str], min_share: float = 0.6) -> list[str]:
    """Remove headers/footers: lines that appear on most pages, plus "Page N" lines.

    Without this, every chunk would contain "DL301 - Neural Network Basics  Page 3", which
    adds noise to embeddings and wastes tokens.
    """
    page_number = re.compile(r"^\s*(page\s*)?\d+(\s*(of|/)\s*\d+)?\s*$", re.IGNORECASE)
    if len(pages) < 2:
        return ["\n".join(line for line in p.splitlines() if not page_number.match(line)) for p in pages]

    def normalise(line: str) -> str:
        return re.sub(r"\d+", "#", line.strip().lower())  # "Page 3" and "Page 4" count as the same line

    counts = Counter(norm for p in pages for norm in {normalise(line) for line in p.splitlines() if line.strip()})
    repeated = {line for line, n in counts.items() if n / len(pages) >= min_share}
    return [
        "\n".join(line for line in p.splitlines() if normalise(line) not in repeated and not page_number.match(line))
        for p in pages
    ]


# ------------------------------------------------------------------------ DOCX


def parse_docx(path: Path) -> list[Section]:
    from docx import Document

    doc = Document(str(path))
    sections: list[Section] = []
    heading, buffer = None, []

    def flush():
        if buffer:
            sections.append(Section(text=_clean("\n".join(buffer)), heading=heading))

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower() if para.style is not None else ""
        if style.startswith("heading") or style == "title":
            flush()
            heading, buffer = text, []
        else:
            buffer.append(f"- {text}" if "list" in style else text)
    flush()
    # Tables are common in course notes: keep each row as one line of text.
    for table in doc.tables:
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        sections.append(Section(text=_clean("\n".join(rows)), heading="Table"))
    return sections


# ------------------------------------------------------------------------ HTML


def parse_html(path: Path) -> list[Section]:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="replace"), "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "header", "noscript"]):
        tag.decompose()  # boilerplate, not course content
    body = soup.body or soup

    sections: list[Section] = []
    heading, buffer = None, []
    for el in body.find_all(["h1", "h2", "h3", "p", "li", "pre", "td"]):
        text = el.get_text(" ", strip=True)
        if not text:
            continue
        if el.name in ("h1", "h2", "h3"):
            if buffer:
                sections.append(Section(text=_clean("\n".join(buffer)), heading=heading))
            heading, buffer = text, []
        else:
            buffer.append(f"- {text}" if el.name == "li" else text)
    if buffer:
        sections.append(Section(text=_clean("\n".join(buffer)), heading=heading))
    return sections


# ------------------------------------------------------------------------- TXT


def parse_txt(path: Path) -> list[Section]:
    """Plain text has no structure, but a line underlined with ---- or ==== is a heading."""
    raw = path.read_bytes()
    try:
        text = raw.decode("utf-8")
    except UnicodeDecodeError:
        text = raw.decode("cp1252", errors="replace")  # common for files saved on Windows

    lines = text.splitlines()
    sections: list[Section] = []
    heading, buffer = None, []
    i = 0
    while i < len(lines):
        line = lines[i]
        underline = lines[i + 1] if i + 1 < len(lines) else ""
        if line.strip() and re.fullmatch(r"[-=]{3,}", underline.strip()):
            if buffer:
                sections.append(Section(text=_clean("\n".join(buffer)), heading=heading))
            heading, buffer = line.strip(), []
            i += 2
            continue
        buffer.append(line)
        i += 1
    if buffer:
        sections.append(Section(text=_clean("\n".join(buffer)), heading=heading))
    return sections
