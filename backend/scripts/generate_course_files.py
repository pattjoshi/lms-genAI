"""Render the course content in data/course_source/*.md into real course files.

    uv run python -m scripts.generate_course_files

Each module is written in a different format so the upload pipeline gets all four
types a teacher might upload:

    module 1 -> PDF   (every topic starts on a new page -> page-number citations)
    module 2 -> DOCX
    module 3 -> HTML
    module 4 -> TXT

Output goes to data/course_files/ and is committed, so you don't need to run this
unless you edit the sources. Needs the dev dependency `reportlab`.
"""

import html
import re
from dataclasses import dataclass, field
from pathlib import Path

from app.config import REPO_ROOT
from app.seed.catalog import COURSES

SOURCE_DIR = REPO_ROOT / "data" / "course_source"
OUT_DIR = REPO_ROOT / "data" / "course_files"
FORMAT_BY_MODULE = {1: "pdf", 2: "docx", 3: "html", 4: "txt"}


@dataclass
class Section:
    title: str
    blocks: list[tuple[str, str | list[str]]] = field(default_factory=list)  # ("p", text) | ("ul", [items])


@dataclass
class Doc:
    title: str
    intro: list[str]
    sections: list[Section]


def parse_markdown(text: str) -> Doc:
    """Tiny parser for our own simple markdown: '# ', '## ', paragraphs and '- ' bullets."""
    title, intro, sections = "", [], []
    current: list = intro
    paragraph: list[str] = []
    bullets: list[str] = []

    def flush():
        nonlocal paragraph, bullets
        if paragraph:
            text_ = " ".join(paragraph)
            current.append(("p", text_) if current is not intro else text_)
            paragraph = []
        if bullets:
            current.append(("ul", bullets))
            bullets = []

    for line in text.splitlines():
        if line.startswith("# "):
            title = line[2:].strip()
        elif line.startswith("## "):
            flush()
            sections.append(Section(line[3:].strip()))
            current = sections[-1].blocks
        elif line.startswith("- "):
            if paragraph:
                flush()
            bullets.append(line[2:].strip())
        elif not line.strip():
            flush()
        else:
            paragraph.append(line.strip())
    flush()
    return Doc(title, intro, sections)


def slug(text: str) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "_", text).strip("_")


# ---------------------------------------------------------------- writers


def write_txt(doc: Doc, path: Path, code: str) -> None:
    lines = [f"{code} - {doc.title}", "=" * (len(code) + 3 + len(doc.title)), ""]
    lines += [p for p in doc.intro] + [""]
    for s in doc.sections:
        lines += [s.title, "-" * len(s.title), ""]
        for kind, value in s.blocks:
            lines += [f"* {item}" for item in value] if kind == "ul" else [value]
            lines.append("")
    path.write_text("\n".join(lines).replace("`", ""), encoding="utf-8")


def _inline_html(text: str) -> str:
    return re.sub(r"`([^`]+)`", r"<code>\1</code>", html.escape(text, quote=False))


def write_html(doc: Doc, path: Path, code: str) -> None:
    parts = [
        "<!doctype html>",
        '<html lang="en"><head><meta charset="utf-8">',
        f"<title>{code} - {html.escape(doc.title)}</title>",
        "<style>body{font-family:Georgia,serif;max-width:46rem;margin:2rem auto;line-height:1.6}"
        "code{background:#f3f3f3;padding:0 .2em}nav{font-size:.9em;color:#666}</style>",
        "</head><body>",
        f"<nav>Course {code} &middot; course notes</nav>",
        f"<h1>{html.escape(doc.title)}</h1>",
    ]
    parts += [f"<p>{_inline_html(p)}</p>" for p in doc.intro]
    for s in doc.sections:
        parts.append(f"<section><h2>{html.escape(s.title)}</h2>")
        for kind, value in s.blocks:
            if kind == "ul":
                parts.append("<ul>" + "".join(f"<li>{_inline_html(i)}</li>" for i in value) + "</ul>")
            else:
                parts.append(f"<p>{_inline_html(value)}</p>")
        parts.append("</section>")
    parts.append("<footer><p>&copy; LMS GenAI demo content</p></footer></body></html>")
    path.write_text("\n".join(parts), encoding="utf-8")


def write_docx(doc: Doc, path: Path, code: str) -> None:
    from docx import Document

    d = Document()
    d.core_properties.title = f"{code} - {doc.title}"
    d.add_heading(doc.title, level=1)
    for p in doc.intro:
        d.add_paragraph(p.replace("`", ""))
    for s in doc.sections:
        d.add_heading(s.title, level=2)
        for kind, value in s.blocks:
            if kind == "ul":
                for item in value:
                    d.add_paragraph(item.replace("`", ""), style="List Bullet")
            else:
                d.add_paragraph(value.replace("`", ""))
    d.save(path)


def write_pdf(doc: Doc, path: Path, code: str) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import ListFlowable, ListItem, PageBreak, Paragraph, SimpleDocTemplate, Spacer

    styles = getSampleStyleSheet()

    def inline(text: str) -> str:
        escaped = html.escape(text, quote=False)
        return re.sub(r"`([^`]+)`", r'<font name="Courier">\1</font>', escaped)

    story = [Paragraph(inline(doc.title), styles["Title"])]
    story += [Paragraph(inline(p), styles["BodyText"]) for p in doc.intro]
    for i, s in enumerate(doc.sections):
        if i > 0:
            story.append(PageBreak())  # one topic per page -> clean page citations
        story += [Spacer(1, 0.3 * cm), Paragraph(inline(s.title), styles["Heading2"])]
        for kind, value in s.blocks:
            if kind == "ul":
                story.append(
                    ListFlowable(
                        [ListItem(Paragraph(inline(item), styles["BodyText"])) for item in value],
                        bulletType="bullet",
                    )
                )
            else:
                story.append(Paragraph(inline(value), styles["BodyText"]))

    def footer(canvas, pdf):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.drawString(2 * cm, 1.2 * cm, f"{code} - {doc.title}")
        canvas.drawRightString(A4[0] - 2 * cm, 1.2 * cm, f"Page {pdf.page}")
        canvas.restoreState()

    pdf = SimpleDocTemplate(str(path), pagesize=A4, title=f"{code} - {doc.title}", author="LMS GenAI demo")
    pdf.build(story, onFirstPage=footer, onLaterPages=footer)


WRITERS = {"pdf": write_pdf, "docx": write_docx, "html": write_html, "txt": write_txt}


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for course in COURSES:
        code = course["code"]
        for number, (module_title, _) in enumerate(course["modules"], start=1):
            doc = parse_markdown((SOURCE_DIR / code / f"m{number}.md").read_text(encoding="utf-8"))
            ext = FORMAT_BY_MODULE[number]
            out = OUT_DIR / f"{code}_M{number}_{slug(module_title)}.{ext}"
            WRITERS[ext](doc, out, code)
            print(f"wrote {out.relative_to(REPO_ROOT)}")


if __name__ == "__main__":
    main()
